"""Integration smoke tests for installed runtime dependencies."""

from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path

from pdf_excel.converter import convert_pdf_to_excel


def _has_runtime_dependencies() -> bool:
    return (
        importlib.util.find_spec("openpyxl") is not None
        and importlib.util.find_spec("pdfplumber") is not None
    )


def _escape_pdf_text(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _simple_text_pdf(lines: list[str]) -> bytes:
    text_lines = ["BT", "/F1 12 Tf", "72 720 Td"]
    for index, line in enumerate(lines):
        if index:
            text_lines.append("0 -18 Td")
        text_lines.append(f"({_escape_pdf_text(line)}) Tj")
    text_lines.append("ET")

    stream = "\n".join(text_lines).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
    ]

    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode("ascii"))
        pdf.extend(obj)
        pdf.extend(b"\nendobj\n")

    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(pdf)


@unittest.skipUnless(_has_runtime_dependencies(), "runtime dependencies not installed")
class ConversionIntegrationTests(unittest.TestCase):
    def test_convert_text_pdf_to_xlsx_workbook(self) -> None:
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as tmpdir:
            pdf_path = Path(tmpdir) / "sample.pdf"
            output_path = Path(tmpdir) / "sample.xlsx"
            pdf_path.write_bytes(
                _simple_text_pdf(["Name Amount", "Alice 10", "Bob 20"])
            )

            result = convert_pdf_to_excel(pdf_path, output_path=output_path)

            self.assertEqual(result.page_count, 1)
            self.assertEqual(result.sheet_count, 1)
            self.assertEqual(result.row_count, 3)
            self.assertTrue(output_path.exists())

            workbook = load_workbook(output_path)
            worksheet = workbook.active
            values = [row[0] for row in worksheet.iter_rows(values_only=True)]
            self.assertEqual(values, ["Name Amount", "Alice 10", "Bob 20"])


if __name__ == "__main__":
    unittest.main()
