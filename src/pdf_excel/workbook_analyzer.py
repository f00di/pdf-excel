"""Analyze Excel workbooks before appending extracted transactions."""

from __future__ import annotations

from collections import Counter
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from pdf_excel.models import SheetStructure, WorkbookAnalysis


HEADER_KEYWORDS = {
    "date",
    "transaction",
    "description",
    "details",
    "narration",
    "particulars",
    "debit",
    "withdrawal",
    "withdrawals",
    "credit",
    "deposit",
    "deposits",
    "balance",
    "amount",
    "category",
    "notes",
    "reference",
    "ref",
}

DATE_HEADER_NAMES = {
    "date",
    "transaction date",
    "trans date",
    "posting date",
    "value date",
}


class ExcelAnalysisError(Exception):
    """Raised when an Excel file cannot be analyzed safely."""


def analyze_excel_path(
    workbook_path: str | Path,
    *,
    password: str | None = None,
) -> WorkbookAnalysis:
    """Analyze an Excel workbook from disk."""

    path = Path(workbook_path).expanduser().resolve()
    return analyze_excel_file(
        file_bytes=path.read_bytes(),
        file_name=path.name,
        password=password,
    )


def analyze_excel_file(
    file_bytes: bytes,
    file_name: str,
    password: str | None = None,
) -> WorkbookAnalysis:
    """Detect sheet headers, formulas, empty areas, and append rows."""

    suffix = Path(file_name).suffix.lower()
    if suffix == ".xls":
        return _analyze_legacy_xls(file_bytes, file_name)

    if suffix not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        raise ExcelAnalysisError(
            "Unsupported Excel file type. Use .xlsx or .xlsm for append output."
        )

    try:
        workbook = load_workbook(BytesIO(file_bytes), data_only=False)
    except Exception as exc:
        message = str(exc)
        if password:
            decrypted = _decrypt_excel_file(file_bytes, password)
            if decrypted:
                workbook = load_workbook(BytesIO(decrypted), data_only=False)
            else:
                raise ExcelAnalysisError(
                    "Could not decrypt this workbook with the supplied password."
                ) from exc
        else:
            if "password" in message.lower() or "encrypted" in message.lower():
                raise ExcelAnalysisError(
                    "This workbook appears to be protected or encrypted. Enter "
                    "the password or upload an unlocked copy."
                ) from exc
            raise ExcelAnalysisError(f"Could not read workbook: {message}") from exc

    sheets = {
        worksheet.title: analyze_worksheet(worksheet)
        for worksheet in workbook.worksheets
    }
    warnings = []
    if suffix in {".xlsm", ".xltm"}:
        warnings.append(
            "Macro-enabled workbooks can be read, but macros are not executed. "
            "Save a backup before using the output."
        )

    return WorkbookAnalysis(
        file_name=file_name,
        file_type=suffix.lstrip("."),
        sheets=sheets,
        writable=True,
        warnings=warnings,
    )


def _decrypt_excel_file(file_bytes: bytes, password: str) -> bytes | None:
    try:
        import msoffcrypto
    except ImportError:
        return None

    try:
        office_file = msoffcrypto.OfficeFile(BytesIO(file_bytes))
        office_file.load_key(password=password)
        decrypted = BytesIO()
        office_file.decrypt(decrypted)
        return decrypted.getvalue()
    except Exception:
        return None


def _analyze_legacy_xls(file_bytes: bytes, file_name: str) -> WorkbookAnalysis:
    try:
        import pandas as pd
    except ImportError as exc:
        raise ExcelAnalysisError(
            ".xls preview requires pandas and xlrd. Save the file as .xlsx "
            "before appending."
        ) from exc

    try:
        excel_file = pd.ExcelFile(BytesIO(file_bytes))
    except Exception as exc:
        raise ExcelAnalysisError(
            f"Could not read .xls workbook. Save it as .xlsx if possible. Details: {exc}"
        ) from exc

    sheets: dict[str, SheetStructure] = {}
    for sheet_name in excel_file.sheet_names:
        frame = excel_file.parse(sheet_name=sheet_name, header=None, nrows=30)
        header_row = _detect_header_row_from_values(frame.values.tolist())
        headers: dict[str, str] = {}
        if not frame.empty:
            values = frame.iloc[max(header_row - 1, 0)].tolist()
            for index, value in enumerate(values, start=1):
                header = _clean_header(value) or f"Column {get_column_letter(index)}"
                headers[get_column_letter(index)] = header

        preview_rows = []
        for row_index, row in frame.head(20).iterrows():
            row_payload = {"__row__": int(row_index) + 1}
            for col_index, value in enumerate(row.tolist(), start=1):
                row_payload[get_column_letter(col_index)] = _display_value(value)
            preview_rows.append(row_payload)

        sheets[sheet_name] = SheetStructure(
            sheet_name=sheet_name,
            max_row=int(frame.shape[0]),
            max_column=int(frame.shape[1]),
            header_row=header_row,
            headers=headers,
            preview_rows=preview_rows,
            empty_areas=[],
            tables=[],
            merged_cells=[],
            data_validations=[],
            formulas=[],
            column_formats={},
            column_widths={},
            next_row=int(frame.shape[0]) + 1,
            warnings=[
                ".xls files can be previewed only. Save as .xlsx before appending."
            ],
        )

    return WorkbookAnalysis(
        file_name=file_name,
        file_type="xls",
        sheets=sheets,
        writable=False,
        warnings=[
            ".xls is a legacy format. This app does not write to .xls because "
            "formatting preservation requires .xlsx or .xlsm."
        ],
    )


def analyze_worksheet(worksheet) -> SheetStructure:
    """Analyze one openpyxl worksheet."""

    header_row = detect_header_row(worksheet)
    max_row = worksheet.max_row or 1
    max_column = worksheet.max_column or 1
    headers = extract_headers(worksheet, header_row)
    preview_rows = preview_sheet_rows(worksheet, max_rows=30)
    empty_areas = detect_empty_areas(worksheet, header_row)
    tables = extract_tables(worksheet)
    merged_cells = [str(cell_range) for cell_range in worksheet.merged_cells.ranges]
    data_validations = [
        f"{validation.type or 'any'} on {validation.sqref}"
        for validation in worksheet.data_validations.dataValidation
    ]
    formulas = extract_formulas(worksheet)
    column_formats = detect_column_formats(worksheet, header_row)
    column_widths = {
        get_column_letter(index): worksheet.column_dimensions[
            get_column_letter(index)
        ].width
        or 8.43
        for index in range(1, max_column + 1)
    }
    next_row = find_next_append_row(worksheet, header_row)
    warnings = []

    if worksheet.protection.sheet:
        warnings.append(
            "This sheet is protected. Writing may fail unless protection is removed."
        )
    if not headers:
        warnings.append("No clear header row was detected.")

    return SheetStructure(
        sheet_name=worksheet.title,
        max_row=max_row,
        max_column=max_column,
        header_row=header_row,
        headers=headers,
        preview_rows=preview_rows,
        empty_areas=empty_areas,
        tables=tables,
        merged_cells=merged_cells,
        data_validations=data_validations,
        formulas=formulas,
        column_formats=column_formats,
        column_widths=column_widths,
        next_row=next_row,
        warnings=warnings,
    )


def detect_header_row(worksheet, scan_rows: int = 25) -> int:
    rows = list(
        worksheet.iter_rows(
            min_row=1,
            max_row=min(worksheet.max_row or 1, scan_rows),
            values_only=True,
        )
    )
    return _detect_header_row_from_values(rows)


def _detect_header_row_from_values(rows: list[list[Any]]) -> int:
    best_row = 1
    best_score = -1.0
    for index, row in enumerate(rows, start=1):
        values = [_clean_header(value) for value in row]
        non_empty = [value for value in values if value]
        if not non_empty:
            continue

        keyword_hits = sum(
            1
            for value in non_empty
            if any(keyword in value.lower() for keyword in HEADER_KEYWORDS)
        )
        textish = sum(1 for value in non_empty if not _looks_numeric(value))
        unique_ratio = len(set(non_empty)) / max(len(non_empty), 1)
        next_row = rows[index] if index < len(rows) else []
        next_data_score = sum(1 for value in next_row if value not in (None, ""))

        score = (
            len(non_empty) * 1.0
            + keyword_hits * 3.0
            + textish * 0.5
            + unique_ratio
            + min(next_data_score, len(non_empty)) * 0.35
        )
        if score > best_score:
            best_score = score
            best_row = index
    return best_row


def extract_headers(worksheet, header_row: int) -> dict[str, str]:
    headers: dict[str, str] = {}
    explicit_headers: dict[str, str] = {}
    for cell in worksheet[header_row]:
        if cell.column > worksheet.max_column:
            continue
        column_letter = get_column_letter(cell.column)
        header = _clean_header(cell.value)
        if header:
            explicit_headers[column_letter] = header

    if len(explicit_headers) >= 2:
        return _primary_header_block(explicit_headers)

    for cell in worksheet[header_row]:
        column_letter = get_column_letter(cell.column)
        header = _clean_header(cell.value)
        if header:
            headers[column_letter] = header
        elif any(
            worksheet.cell(row=row, column=cell.column).value not in (None, "")
            for row in range(header_row + 1, min(worksheet.max_row + 1, header_row + 8))
        ):
            headers[column_letter] = f"Column {column_letter}"
    return headers


def preview_sheet_rows(worksheet, max_rows: int = 30) -> list[dict[str, Any]]:
    preview = []
    max_column = min(worksheet.max_column or 1, 30)
    for row in worksheet.iter_rows(
        min_row=1,
        max_row=min(worksheet.max_row or 1, max_rows),
        max_col=max_column,
    ):
        payload: dict[str, Any] = {"__row__": row[0].row if row else None}
        for cell in row:
            payload[get_column_letter(cell.column)] = _display_value(cell.value)
        preview.append(payload)
    return preview


def detect_empty_areas(worksheet, header_row: int) -> list[str]:
    empty_ranges = []
    current_start: int | None = None
    headers = extract_headers(worksheet, header_row)
    date_column = _date_column(headers)
    formula_columns = detect_formula_columns(worksheet, header_row)
    entry_columns = _entry_columns(headers, formula_columns)
    for row_number in range(header_row + 1, (worksheet.max_row or 1) + 1):
        row_is_empty = (
            _date_cell_is_empty(worksheet, row_number, date_column)
            if date_column
            else not _row_has_entry_data(worksheet, row_number, entry_columns)
        )
        if row_is_empty and current_start is None:
            current_start = row_number
        if not row_is_empty and current_start is not None:
            empty_ranges.append(f"Rows {current_start}-{row_number - 1}")
            current_start = None
    if current_start is not None:
        empty_ranges.append(f"Rows {current_start}-{worksheet.max_row}")
    return empty_ranges


def extract_tables(worksheet) -> list[dict[str, Any]]:
    tables = []
    for table in worksheet.tables.values():
        tables.append(
            {
                "name": table.name,
                "display_name": table.displayName,
                "ref": table.ref,
                "table_style": table.tableStyleInfo.name
                if table.tableStyleInfo
                else "",
            }
        )
    return tables


def extract_formulas(worksheet, limit: int = 100) -> list[dict[str, str]]:
    formulas = []
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                formulas.append({"cell": cell.coordinate, "formula": str(cell.value)})
                if len(formulas) >= limit:
                    return formulas
    return formulas


def detect_column_formats(worksheet, header_row: int) -> dict[str, dict[str, Any]]:
    formats: dict[str, dict[str, Any]] = {}
    max_row = min(worksheet.max_row or header_row, header_row + 100)
    for column in range(1, (worksheet.max_column or 1) + 1):
        column_letter = get_column_letter(column)
        number_formats = Counter()
        data_types = Counter()
        examples = []
        formula_examples = []
        formula_count = 0
        for row in range(header_row + 1, max_row + 1):
            cell = worksheet.cell(row=row, column=column)
            if cell.value in (None, ""):
                continue
            if cell.data_type == "f":
                formula_count += 1
                if len(formula_examples) < 3:
                    formula_examples.append(str(cell.value))
            number_formats[cell.number_format] += 1
            data_types[type(cell.value).__name__] += 1
            if len(examples) < 3:
                examples.append(_display_value(cell.value))
        formats[column_letter] = {
            "header": _clean_header(
                worksheet.cell(row=header_row, column=column).value
            ),
            "number_format": number_formats.most_common(1)[0][0]
            if number_formats
            else "General",
            "data_type": data_types.most_common(1)[0][0] if data_types else "empty",
            "examples": examples,
            "formula_count": formula_count,
            "formula_examples": formula_examples,
        }
    return formats


def find_next_append_row(worksheet, header_row: int) -> int:
    """Find the next row below real entry data, ignoring formula templates."""

    headers = extract_headers(worksheet, header_row)
    date_column = _date_column(headers)
    if date_column:
        last_date_row = _last_filled_date_row(worksheet, header_row, date_column)
        if last_date_row > header_row:
            return last_date_row + 1

    formula_columns = detect_formula_columns(worksheet, header_row)
    entry_columns = _entry_columns(headers, formula_columns)
    last_data_row = header_row
    for row_number in range(header_row + 1, (worksheet.max_row or header_row) + 1):
        if _row_has_entry_data(worksheet, row_number, entry_columns):
            last_data_row = row_number
    return last_data_row + 1


def detect_formula_columns(worksheet, header_row: int) -> set[str]:
    formula_columns: set[str] = set()
    max_row = worksheet.max_row or header_row
    for column in range(1, (worksheet.max_column or 1) + 1):
        formula_count = 0
        value_count = 0
        for row in range(header_row + 1, max_row + 1):
            cell = worksheet.cell(row=row, column=column)
            if cell.value in (None, ""):
                continue
            value_count += 1
            if cell.data_type == "f":
                formula_count += 1
        if formula_count >= 3 and formula_count / max(value_count, 1) >= 0.5:
            formula_columns.add(get_column_letter(column))
    return formula_columns


def column_to_index(column_letter: str) -> int:
    index = 0
    for char in column_letter.upper():
        if not char.isalpha():
            continue
        index = index * 26 + ord(char) - ord("A") + 1
    return index


def _date_column(headers: dict[str, str]) -> str:
    best_column = ""
    best_score = 0
    for column, header in headers.items():
        normalized = _normalize_header(header)
        if normalized == "date":
            return column
        score = 0
        if normalized in DATE_HEADER_NAMES:
            score = 3
        elif "date" in normalized and "statement" not in normalized:
            score = 1
        if score > best_score:
            best_column = column
            best_score = score
    return best_column if best_score else ""


def _last_filled_date_row(worksheet, header_row: int, date_column: str) -> int:
    column_index = column_to_index(date_column)
    last_date_row = header_row
    for row_number in range(header_row + 1, (worksheet.max_row or header_row) + 1):
        cell = worksheet.cell(row=row_number, column=column_index)
        if cell.value not in (None, "") and cell.data_type != "f":
            last_date_row = row_number
    return last_date_row


def _date_cell_is_empty(worksheet, row_number: int, date_column: str) -> bool:
    column_index = column_to_index(date_column)
    cell = worksheet.cell(row=row_number, column=column_index)
    return cell.value in (None, "") or cell.data_type == "f"


def _primary_header_block(headers: dict[str, str]) -> dict[str, str]:
    if not headers:
        return {}

    ordered_columns = sorted(headers, key=lambda column: column_to_index(column))
    primary: dict[str, str] = {}
    previous_index: int | None = None
    for column in ordered_columns:
        column_index = column_to_index(column)
        if previous_index is not None and column_index > previous_index + 1:
            break
        primary[column] = headers[column]
        previous_index = column_index
    return primary


def _entry_columns(headers: dict[str, str], formula_columns: set[str]) -> list[int]:
    columns = []
    for column_letter in headers:
        if column_letter in formula_columns:
            continue
        columns.append(column_to_index(column_letter))
    return columns


def _row_has_entry_data(worksheet, row_number: int, entry_columns: list[int]) -> bool:
    for column in entry_columns:
        cell = worksheet.cell(row=row_number, column=column)
        if cell.value in (None, ""):
            continue
        if cell.data_type == "f":
            continue
        return True
    return False


def _normalize_header(value: Any) -> str:
    return " ".join(str(value or "").replace("\n", " ").lower().split()).strip()


def _clean_header(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\n", " ").split()).strip()


def _display_value(value: Any) -> Any:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def _looks_numeric(value: Any) -> bool:
    try:
        float(str(value).replace(",", "").replace("$", ""))
        return True
    except ValueError:
        return False

