"""Write extracted transactions into existing Excel workbooks."""

from __future__ import annotations

from copy import copy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from pdf_excel.models import AppendOperation, ExtractedTransaction, WriteSummary
from pdf_excel.statement_extractor import normalize_date
from pdf_excel.workbook_analyzer import find_next_append_row


HIGHLIGHT_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
LOG_SHEET_NAME = "_PDF_Append_Log"


class ExcelWriteError(Exception):
    """Raised when a workbook cannot be written safely."""


def append_transactions_to_workbook(
    excel_bytes: bytes,
    source_excel_name: str,
    transactions: list[ExtractedTransaction],
    operations: list[AppendOperation],
    *,
    password: str | None = None,
) -> tuple[bytes, bytes, WriteSummary]:
    """Append transactions into an existing workbook and return output bytes."""

    suffix = Path(source_excel_name).suffix.lower()
    if suffix == ".xls":
        raise ExcelWriteError(
            "Writing to .xls is not supported. Save the workbook as .xlsx first."
        )

    try:
        workbook = load_workbook(
            BytesIO(excel_bytes),
            data_only=False,
            keep_vba=suffix in {".xlsm", ".xltm"},
        )
    except Exception as exc:
        decrypted = _decrypt_excel_file(excel_bytes, password) if password else None
        if not decrypted:
            raise ExcelWriteError(f"Could not open workbook for writing: {exc}") from exc
        workbook = load_workbook(
            BytesIO(decrypted),
            data_only=False,
            keep_vba=suffix in {".xlsm", ".xltm"},
        )

    changed_cells: list[dict[str, Any]] = []
    warnings: list[str] = []
    rows_added = 0
    rows_skipped = 0

    log_sheet = _ensure_log_sheet(workbook)
    for operation in operations:
        if operation.target_sheet not in workbook.sheetnames:
            warnings.append(f"Target sheet not found: {operation.target_sheet}.")
            continue

        worksheet = workbook[operation.target_sheet]
        if worksheet.protection.sheet:
            warnings.append(
                f"Sheet {operation.target_sheet} is protected; attempted write "
                "may be rejected by Excel."
            )

        start_row = _resolve_start_row(worksheet, operation)
        existing_keys = (
            _existing_keys(worksheet, operation)
            if operation.write_mode == "skip_duplicates"
            else set()
        )

        for transaction in transactions:
            transaction_key = _transaction_key(transaction)
            if operation.write_mode == "skip_duplicates" and transaction_key in existing_keys:
                rows_skipped += 1
                continue

            target_row = start_row
            if operation.write_mode != "overwrite":
                while _row_has_values(
                    worksheet,
                    target_row,
                    operation.field_mapping.values(),
                ):
                    target_row += 1

            _copy_row_format_and_formulas(
                worksheet=worksheet,
                source_row=max(operation.header_row + 1, target_row - 1),
                target_row=target_row,
                mapped_columns=set(operation.field_mapping.values()),
            )

            wrote_anything = False
            for field, column_letter in operation.field_mapping.items():
                value = _transaction_field_value(transaction, field)
                if value in (None, ""):
                    continue

                cell = worksheet[f"{column_letter}{target_row}"]
                if _is_merged_non_anchor(worksheet, cell.coordinate):
                    warnings.append(f"Skipped merged non-anchor cell {cell.coordinate}.")
                    continue

                cell.value = _coerce_for_excel(field, value)
                if operation.highlight_new_cells:
                    cell.fill = copy(HIGHLIGHT_FILL)

                changed_cell = {
                    "sheet": operation.target_sheet,
                    "cell": cell.coordinate,
                    "field": field,
                    "value": value,
                    "source_pdf": transaction.source_pdf,
                }
                changed_cells.append(changed_cell)
                _append_log_row(log_sheet, changed_cell, target_row)
                wrote_anything = True

            if wrote_anything:
                rows_added += 1
                existing_keys.add(transaction_key)
                start_row = target_row + 1
            else:
                rows_skipped += 1

    output = BytesIO()
    workbook.save(output)
    output_bytes = output.getvalue()
    output_name = _updated_file_name(source_excel_name)

    summary = WriteSummary(
        source_excel_name=source_excel_name,
        target_sheets=[operation.target_sheet for operation in operations],
        rows_added=rows_added,
        rows_skipped=rows_skipped,
        output_file_name=output_name,
        warnings=warnings,
        changed_cells=changed_cells,
        backup_created=True,
    )
    return output_bytes, excel_bytes, summary


def _decrypt_excel_file(file_bytes: bytes, password: str | None) -> bytes | None:
    if not password:
        return None
    try:
        import msoffcrypto
    except ImportError:
        return None

    try:
        office_file = msoffcrypto.OfficeFile(BytesIO(file_bytes))
        office_file.load_key(password=password)
        decrypted = BytesIO()
        office_file.decrypt(decrypted)
        return decrypted.getvalue()
    except Exception:
        return None


def _resolve_start_row(worksheet, operation: AppendOperation) -> int:
    if operation.write_mode == "overwrite":
        return max(operation.start_row, operation.header_row + 1)
    next_append_row = find_next_append_row(worksheet, operation.header_row)
    return max(operation.start_row, next_append_row)


def _copy_row_format_and_formulas(
    worksheet,
    source_row: int,
    target_row: int,
    mapped_columns: set[str],
) -> None:
    if source_row < 1 or source_row == target_row:
        return

    worksheet.row_dimensions[target_row].height = worksheet.row_dimensions[
        source_row
    ].height
    max_column = worksheet.max_column or 1
    for column_index in range(1, max_column + 1):
        column_letter = get_column_letter(column_index)
        source_cell = worksheet.cell(row=source_row, column=column_index)
        target_cell = worksheet.cell(row=target_row, column=column_index)

        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
        if source_cell.number_format:
            target_cell.number_format = source_cell.number_format
        if source_cell.alignment:
            target_cell.alignment = copy(source_cell.alignment)
        if source_cell.protection:
            target_cell.protection = copy(source_cell.protection)

        if column_letter in mapped_columns:
            continue
        if source_cell.data_type == "f" and source_cell.value:
            try:
                target_cell.value = Translator(
                    source_cell.value,
                    origin=source_cell.coordinate,
                ).translate_formula(target_cell.coordinate)
            except Exception:
                target_cell.value = source_cell.value


def _row_has_values(worksheet, row: int, columns: Any) -> bool:
    for column_letter in columns:
        cell = worksheet[f"{column_letter}{row}"]
        if cell.data_type == "f":
            continue
        if cell.value not in (None, ""):
            return True
    return False


def _existing_keys(worksheet, operation: AppendOperation) -> set[str]:
    keys = set()
    date_col = operation.field_mapping.get("transaction_date")
    desc_col = operation.field_mapping.get("description")
    debit_col = operation.field_mapping.get("debit")
    credit_col = operation.field_mapping.get("credit")
    if not date_col or not desc_col:
        return keys

    for row in range(operation.header_row + 1, (worksheet.max_row or operation.header_row) + 1):
        key = _row_key(
            worksheet[f"{date_col}{row}"].value,
            worksheet[f"{desc_col}{row}"].value,
            worksheet[f"{debit_col}{row}"].value if debit_col else None,
            worksheet[f"{credit_col}{row}"].value if credit_col else None,
        )
        if key:
            keys.add(key)
    return keys


def _ensure_log_sheet(workbook):
    if LOG_SHEET_NAME in workbook.sheetnames:
        worksheet = workbook[LOG_SHEET_NAME]
    else:
        worksheet = workbook.create_sheet(LOG_SHEET_NAME)
        worksheet.append(
            ["Timestamp", "Source PDF", "Target Sheet", "Target Row", "Cell", "Field", "Value"]
        )
    return worksheet


def _append_log_row(log_sheet, changed_cell: dict[str, Any], target_row: int) -> None:
    log_sheet.append(
        [
            datetime.now().isoformat(timespec="seconds"),
            changed_cell.get("source_pdf", ""),
            changed_cell.get("sheet", ""),
            target_row,
            changed_cell.get("cell", ""),
            changed_cell.get("field", ""),
            changed_cell.get("value", ""),
        ]
    )


def _transaction_field_value(transaction: ExtractedTransaction, field: str) -> Any:
    return getattr(transaction, field, "")


def _coerce_for_excel(field: str, value: Any) -> Any:
    if field == "transaction_date":
        normalized = normalize_date(value)
        if normalized:
            return date.fromisoformat(normalized)
        return value
    if field in {"debit", "credit", "balance", "opening_balance", "closing_balance"}:
        try:
            return float(value)
        except (TypeError, ValueError):
            return value
    return value


def _is_merged_non_anchor(worksheet, coordinate: str) -> bool:
    for merged_range in worksheet.merged_cells.ranges:
        if coordinate in merged_range:
            return coordinate != merged_range.start_cell.coordinate
    return False


def _transaction_key(transaction: ExtractedTransaction) -> str:
    return _row_key(
        transaction.transaction_date,
        transaction.description,
        transaction.debit,
        transaction.credit,
    )


def _row_key(date_value: Any, description: Any, debit: Any, credit: Any) -> str:
    date_text = normalize_date(date_value) or str(date_value or "").strip()
    description_text = " ".join(str(description or "").lower().split())
    debit_text = _amount_text(debit)
    credit_text = _amount_text(credit)
    if not any([date_text, description_text, debit_text, credit_text]):
        return ""
    return "|".join([date_text, description_text, debit_text, credit_text])


def _amount_text(value: Any) -> str:
    if value in (None, ""):
        return ""
    try:
        return f"{float(value):.2f}"
    except (TypeError, ValueError):
        return str(value).strip()


def _updated_file_name(source_excel_name: str) -> str:
    path = Path(source_excel_name)
    suffix = path.suffix or ".xlsx"
    return f"{path.stem}_updated{suffix}"

